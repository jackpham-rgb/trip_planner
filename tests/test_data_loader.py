from pathlib import Path

import openpyxl
import pytest

from build_option_bank import build_option_bank, parse_sheet

REAL_WORKBOOK = Path(__file__).parent.parent / "Jack-Master-Travel-Database.xlsx"


def _make_synthetic_sheet(wb, name, blocks):
    """blocks: list of (section_title, [data_rows]) using the 27-col header shape."""
    ws = wb.create_sheet(name)
    header = ["Version", "Day", "Date", "Country / State", "City / Region", "Overnight",
              "Route (From -> To)", "Transport", "Transit", "Main Destination",
              "Specific Attraction / Activity", "Address / Area", "Website", "Why It's Worth It",
              "Time Needed", "Priority", "Type", "Time of Day", "Cost", "Reserve?", "Guide?",
              "Seasonal?", "Weather Dep.?", "Viral Rating", "Food / Drink Pick", "Info Class",
              "Best Time / Recommended Date"]
    for title, rows in blocks:
        ws.append([title])
        ws.append([])
        ws.append(header)
        for r in rows:
            ws.append(r)
        ws.append([])
    return ws


def _data_row(dest, hours_text="2h", type_="Nature"):
    row = [None] * 27
    row[0], row[1] = "Solo 2W", 1
    row[3], row[4] = "TestLand", "TestCity"
    row[9] = dest
    row[10] = f"{dest} attraction"
    row[14] = hours_text
    row[16] = type_
    row[15] = "HIGH"
    row[18] = "$$"
    row[23] = "Actually worth it"
    return row


def test_parse_sheet_extracts_data_rows_under_section_title():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = _make_synthetic_sheet(wb, "Synthetic", [
        ("Test Region", [_data_row("Place A"), _data_row("Place B", hours_text="3 days")]),
    ])
    options = parse_sheet(ws, "Synthetic")
    assert len(options) == 2
    assert all(o.section == "Test Region" for o in options)
    assert options[0].main_destination == "Place A"
    assert options[1].est_hours == 24.0  # "3 days" -> 3*8


def test_parse_sheet_handles_multiple_blocks_and_unique_ids():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = _make_synthetic_sheet(wb, "Synthetic", [
        ("Region One", [_data_row("Alpha")]),
        ("Region Two", [_data_row("Alpha")]),  # same destination name, different section
    ])
    options = parse_sheet(ws, "Synthetic")
    assert len(options) == 2
    assert len(set(o.id for o in options)) == 2  # IDs must not collide


def test_parse_sheet_ignores_non_itinerary_overview_tables():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Synthetic")
    ws.append(["OVERVIEW"])
    ws.append(["Field", "Detail"])
    ws.append(["Design rule", "Some text"])
    ws.append([])
    header = ["Version", "Day"] + [None] * 25
    ws.append(header)
    ws.append(_data_row("Real Item"))
    options = parse_sheet(ws, "Synthetic")
    assert len(options) == 1
    assert options[0].main_destination == "Real Item"


@pytest.mark.skipif(not REAL_WORKBOOK.exists(), reason="real workbook not present in this checkout")
def test_build_option_bank_on_real_workbook_produces_options():
    options = build_option_bank(REAL_WORKBOOK)
    assert len(options) > 200
    assert len(set(o.id for o in options)) == len(options)  # all IDs unique
    assert any(o.type == "Adventure" for o in options)
